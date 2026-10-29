from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck_prototype():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_navy = RGBColor(10, 14, 26)       # #0a0e1a
    c_card = RGBColor(15, 23, 42)       # #0f172a
    c_border = RGBColor(30, 41, 59)     # #1e293b
    c_cyan = RGBColor(0, 212, 255)      # #00d4ff
    c_blue = RGBColor(0, 80, 255)       # #0050ff
    c_gold = RGBColor(245, 200, 66)     # #f5c842
    c_white = RGBColor(255, 255, 255)   # #ffffff
    c_text = RGBColor(200, 216, 240)    # #c8d8f0
    c_muted = RGBColor(148, 163, 184)   # #94a3b8

    def add_base_slide():
        s = prs.slides.add_slide(blank_layout)
        # Background
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_navy
        bg.line.fill.background()
        # Top Accent Line
        top_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = c_cyan
        top_bar.line.fill.background()
        # Footer branding
        ft = s.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.4))
        p = ft.text_frame.paragraphs[0]
        p.text = "TERRESTRIAL PLATFORM™ | ENTERPRISE INDUSTRIAL INTELLIGENCE SUITE | CONFIDENTIAL — FOR SYSTEM INTEGRATORS"
        p.font.size = Pt(8.5)
        p.font.color.rgb = c_muted
        return s

    # Slide 1: Cover
    s1 = add_base_slide()
    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(3.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "ENTERPRISE APPLICATION CATALOGUE | AUTONOMOUS AI & SWARMS"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = c_cyan
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "TerraAI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = c_white
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous Industrial Intelligence & Physics-Informed Predictive Operations Engine"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = c_gold
    p2.space_after = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Protocol Standard: OPC UA / TSN (IEEE 802.1Qbv) | IEC 61850 | MQTT Sparkplug B"
    p3.font.size = Pt(13)
    p3.font.color.rgb = c_cyan

    # 4 Metric Boxes on Slide 1
    metrics = [
        ("AUTONOMOUS CONTROL", "Level 4 Autonomy"),
        ("PREDICTIVE HORIZON", "72h Lookahead"),
        ("ANALYTICS ENGINE", "PINO Neural Operators"),
        ("CLOSED-LOOP LATENCY", "Sub-25ms Actuation")
    ]
    card_w = Inches(2.75)
    card_h = Inches(1.3)
    card_gap = Inches(0.24)
    left_start = Inches(0.8)
    top_pos = Inches(4.8)

    for i, (lbl, val) in enumerate(metrics):
        cur_left = left_start + i * (card_w + card_gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_left, top_pos, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp1 = ctf.paragraphs[0]
        cp1.text = lbl
        cp1.alignment = PP_ALIGN.CENTER
        cp1.font.size = Pt(9.5)
        cp1.font.bold = True
        cp1.font.color.rgb = c_muted
        cp1.space_after = Pt(4)

        cp2 = ctf.add_paragraph()
        cp2.text = val
        cp2.alignment = PP_ALIGN.CENTER
        cp2.font.size = Pt(14)
        cp2.font.bold = True
        cp2.font.color.rgb = c_cyan

    prs.save(r'c:\Users\yashj\Downloads\website\test_cover.pptx')
    print("Saved test_cover.pptx successfully!")

create_deck_prototype()

