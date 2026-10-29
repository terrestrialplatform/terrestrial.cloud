import os
import re
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
platform_dir = r'c:\Users\yashj\Downloads\website\terrestrial.cloud\platform'
out_dirs = [
    r'c:\Users\yashj\Downloads\website\si-portal\catalogues',
    r'c:\Users\yashj\Downloads\website\catalogues',
    r'c:\Users\yashj\Downloads\Terrestrial_metamorphosis\si-portal\catalogues'
]

for d in out_dirs:
    os.makedirs(d, exist_ok=True)

# Catalog metadata from platform/index.html
with open(os.path.join(platform_dir, 'index.html'), 'r', encoding='utf-8') as f:
    index_text = f.read()

catalog_data = {}
cards = re.findall(r'<div class="catalog-card".*?</a>\s*</div>', index_text, re.DOTALL)
for c in cards:
    h3_m = re.search(r'<h3>(.*?)</h3>', c)
    if not h3_m:
        continue
    name = h3_m.group(1).strip()
    cat_m = re.search(r'<span class="catalog-cat">(.*?)</span>', c)
    cat = cat_m.group(1).strip() if cat_m else ''
    tag_m = re.search(r'<span class="catalog-protocol-tag">(.*?)</span>', c)
    tag = tag_m.group(1).strip() if tag_m else ''
    brief_m = re.search(r'<p class="catalog-brief">(.*?)</p>', c, re.DOTALL)
    brief = re.sub(r'\s+', ' ', brief_m.group(1)).strip() if brief_m else ''
    catalog_data[name.lower()] = {
        'name': name,
        'category': cat,
        'protocol': tag,
        'brief': brief
    }

# App files list
app_files = [
    "terra-ai.html",
    "terra-cognizance.html",
    "terra-embodiment.html",
    "terra-hive.html",
    "terra-cc.html",
    "terra-scada.html",
    "terra-avatar.html",
    "terra-assets.html",
    "terra-pfd.html",
    "terra-alarm.html",
    "tpim.html",
    "terra-gridx.html",
    "terra-gridw.html",
    "terra-gridg.html",
    "terra-corroguard.html",
    "terra-solar.html",
    "terra-ems.html",
    "terra-3rdeye.html",
    "terra-bms.html",
    "terra-hvac.html"
]

def clean_txt(text):
    if not text:
        return ""
    text = re.sub(r'<[^>]+>', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def parse_app_file(fname):
    fpath = os.path.join(platform_dir, fname)
    with open(fpath, 'r', encoding='utf-8', errors='ignore') as fp:
        html = fp.read()
    
    soup = BeautifulSoup(html, 'html.parser')
    
    # Title / H1
    h1 = soup.find('h1')
    h1_text = clean_txt(h1.get_text()) if h1 else fname
    if ':' in h1_text:
        app_name = h1_text.split(':', 1)[0].strip()
        tagline = h1_text.split(':', 1)[1].strip()
    else:
        app_name = h1_text
        tagline = ''
    
    lookup = catalog_data.get(app_name.lower(), {})
    category = lookup.get('category', 'Industrial Intelligence Suite')
    protocol = lookup.get('protocol', 'OPC UA / TSN / IEC 61850')
    brief = lookup.get('brief', tagline)

    # 1. Telemetry Metrics
    metrics = []
    for card in soup.find_all(class_=lambda c: c and any(k in c for k in ['telemetry-metric', 'metric-card', 'hero-metric', 'spec-item', 'telemetry-item', 'telemetry-pill'])):
        lbl = card.find(class_=lambda c: c and ('label' in c or 'sub' in c))
        val = card.find(class_=lambda c: c and ('value' in c or 'val' in c or 'num' in c))
        if lbl and val:
            metrics.append((clean_txt(lbl.get_text()), clean_txt(val.get_text())))
        else:
            txt = clean_txt(card.get_text())
            if ':' in txt:
                p1, p2 = txt.split(':', 1)
                metrics.append((p1.strip(), p2.strip()))

    # Fallback metrics if few
    if len(metrics) < 4:
        defaults = [
            ("DEPLOYMENT MODEL", "Edge + Hybrid Cloud"),
            ("EXECUTION LATENCY", "Sub-25ms Deterministic"),
            ("INTEROPERABILITY", protocol),
            ("SYSTEM AVAILABILITY", "99.999% Fault-Tolerant")
        ]
        for d in defaults:
            if len(metrics) < 4:
                metrics.append(d)

    # 2. Features / What it will do
    features = []
    sec_features = None
    for s in soup.find_all('section'):
        h2 = s.find(['h2', 'h3'])
        if h2 and 'will do' in h2.get_text().lower():
            sec_features = s
            break
    
    if sec_features:
        for card in sec_features.find_all(class_=lambda c: c and any(k in c for k in ['capability-card', 'feature-card', 'showcase-card', 'col-', 'grid-card', 'card'])):
            h = card.find(['h3', 'h4', 'strong'])
            p = card.find('p')
            if h and p:
                t_str = clean_txt(h.get_text())
                p_str = clean_txt(p.get_text())
                if len(t_str) > 3 and len(p_str) > 15 and not any(f[0] == t_str for f in features):
                    features.append((t_str, p_str))

    # Fallback features if empty
    if len(features) < 4:
        features.append(("Autonomous Real-Time Ingestion", f"High-speed data normalization and continuous telemetry capture across {protocol}."))
        features.append(("Physics-Informed Digital Twin", "Dynamic mathematical modeling mapping physical machine kinematics, thermal limits, and real-time stress envelopes."))
        features.append(("Supervisory Setpoint Optimization", "Continuous closed-loop actuation optimizing throughput, minimizing raw utility consumption, and suppressing process drift."))
        features.append(("Enterprise ERP & EAM Integration", "Automated work order creation, predictive spare parts allocation, and bi-directional SCADA synchronization."))

    # 3. Use cases
    use_cases = []
    sec_uc = None
    for s in soup.find_all('section'):
        h2 = s.find(['h2', 'h3'])
        if h2 and 'use cases' in h2.get_text().lower():
            sec_uc = s
            break
            
    if sec_uc:
        # Check usecase-box
        boxes = sec_uc.find_all(class_=lambda c: c and 'usecase-box' in c)
        if not boxes:
            boxes = sec_uc.find_all(class_=lambda c: c and any(k in c for k in ['card', 'col-', 'item']))
        for b in boxes:
            h = b.find(['h4', 'h3', 'strong'])
            p = b.find('p')
            badge = b.find(class_=lambda c: c and 'badge' in c)
            badge_txt = clean_txt(badge.get_text()) if badge else "Industrial Deployment"
            if h and p:
                t_str = clean_txt(h.get_text())
                p_str = clean_txt(p.get_text())
                if len(t_str) > 3 and len(p_str) > 20 and not any(u[1] == t_str for u in use_cases):
                    use_cases.append((badge_txt, t_str, p_str))

    # Fallback use cases
    if len(use_cases) < 3:
        use_cases.append(("Critical Plant Reliability", f"{app_name} Anomaly Prevention", f"Continuous high-frequency analytics detect sub-synchronous resonance and process drift, eliminating catastrophic asset failure."))
        use_cases.append(("Energy & Utility Optimization", f"{app_name} Efficiency Maximizer", f"Dynamic setpoint modulation reduces specific power and fuel consumption by 2% to 5% across peak operational cycles."))
        use_cases.append(("Autonomous Safe Operations", f"{app_name} Automated Guardrails", f"Real-time safety interlocks and ISA-conforming supervisory actions safeguard plant personnel and critical equipment."))

    # 4. Comparisons
    comparisons = []
    sec_comp = None
    for s in soup.find_all('section'):
        h2 = s.find(['h2', 'h3'])
        if h2 and 'comparison' in h2.get_text().lower():
            sec_comp = s
            break
            
    if sec_comp:
        for tr in sec_comp.find_all('tr'):
            tds = tr.find_all(['td'])
            if len(tds) >= 2:
                c1 = clean_txt(tds[0].get_text())
                c2 = clean_txt(tds[1].get_text())
                if len(c1) > 10 and len(c2) > 10:
                    comparisons.append((c1, c2))

    # Fallback comparisons
    if len(comparisons) < 4:
        comparisons.append(("Static high/low threshold alarms trigger only after physical process degradation has occurred.", f"{app_name} physics-informed models forecast multi-variable trajectories up to 72 hours in advance."))
        comparisons.append(("Operators must manually adjust setpoints across disparate siloed HMIs.", f"Sub-25ms automated closed-loop optimization executes setpoints directly through {protocol}."))
        comparisons.append(("Cryptic tag IDs and fragmented menus cause severe cognitive load during emergencies.", f"Unified semantic digital twin contextualizes all telemetry into actionable plant-wide intelligence."))
        comparisons.append(("Alarm floods of 1,000+ alarms/min overwhelm operators during trips.", "Directed causal graphs de-duplicate alarm avalanches, pinpointing root cause in under 50ms."))

    # 5. Architecture Pipeline
    pipeline = []
    sec_pipe = None
    for s in soup.find_all('section'):
        h2 = s.find(['h2', 'h3'])
        if h2 and any(k in h2.get_text().lower() for k in ['diagram', 'pipeline', 'architecture']):
            sec_pipe = s
            break
            
    if sec_pipe:
        nodes = sec_pipe.find_all(class_=lambda c: c and 'diagram-node' in c)
        for n in nodes:
            t = n.find(class_=lambda c: c and 'node-title' in c)
            d = n.find(class_=lambda c: c and 'node-desc' in c)
            l = n.find(class_=lambda c: c and 'layer-tag' in c)
            layer_txt = clean_txt(l.get_text()) if l else "Architecture Layer"
            if t and d:
                pipeline.append((layer_txt, clean_txt(t.get_text()), clean_txt(d.get_text())))

    if len(pipeline) < 4:
        pipeline = [
            ("Layer 01", "Multi-Protocol Field Telemetry Ingestion", f"Sub-millisecond data acquisition aggregating PLCs, DCS, RTUs, and smart sensor buses via {protocol}."),
            ("Layer 02", "Deterministic Edge Streaming Fabric", "Local real-time microservices execute sub-25ms closed-loop logic and data buffering with zero cloud reliance."),
            ("Layer 03", "Physics-Informed AI Reasoning Engine", "High-performance neural operators compute real-time mass/energy balance, vibration spectra, and predictive trajectories."),
            ("Layer 04", "Supervisory Closed-Loop & Cloud Core", "Multi-site orchestration, cross-plant benchmarking, and bidirectional ERP/SAP maintenance dispatch.")
        ]

    return {
        'name': app_name,
        'tagline': tagline,
        'category': category,
        'protocol': protocol,
        'brief': brief,
        'metrics': metrics[:4],
        'features': features[:4],
        'use_cases': use_cases[:3],
        'comparisons': comparisons[:4],
        'pipeline': pipeline[:4]
    }

# PPTX Generator
def generate_pptx_for_app(data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_navy = RGBColor(10, 14, 26)       # #0a0e1a
    c_card = RGBColor(15, 23, 42)       # #0f172a
    c_card_alt = RGBColor(20, 30, 55)   # #141e37
    c_border = RGBColor(0, 212, 255)    # #00d4ff
    c_border_sub = RGBColor(30, 41, 59) # #1e293b
    c_cyan = RGBColor(0, 212, 255)      # #00d4ff
    c_gold = RGBColor(245, 200, 66)     # #f5c842
    c_green = RGBColor(0, 230, 142)     # #00e68e
    c_white = RGBColor(255, 255, 255)   # #ffffff
    c_text = RGBColor(200, 216, 240)    # #c8d8f0
    c_muted = RGBColor(148, 163, 184)   # #94a3b8
    c_red = RGBColor(255, 85, 85)       # #ff5555

    def add_base_slide(header_tag, title_text):
        s = prs.slides.add_slide(blank_layout)
        # Background
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_navy
        bg.line.fill.background()
        
        # Top Cyan Line
        top_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = c_cyan
        top_bar.line.fill.background()

        if header_tag and title_text:
            tb = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
            tf = tb.text_frame
            tf.word_wrap = True
            
            p0 = tf.paragraphs[0]
            p0.text = f"{header_tag}  ·  TERRESTRIAL PLATFORM™"
            p0.font.size = Pt(10)
            p0.font.bold = True
            p0.font.color.rgb = c_cyan
            p0.space_after = Pt(4)

            p1 = tf.add_paragraph()
            p1.text = title_text
            p1.font.size = Pt(22)
            p1.font.bold = True
            p1.font.color.rgb = c_white

        # Footer branding
        ft = s.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
        p = ft.text_frame.paragraphs[0]
        p.text = f"TERRESTRIAL PLATFORM™  |  {data['name'].upper()} CATALOGUE  |  CONFIDENTIAL — AUTHORISED SYSTEM INTEGRATORS ONLY"
        p.font.size = Pt(8.5)
        p.font.color.rgb = c_muted
        return s

    # ══════════════════════════════════════════════════════════════
    # SLIDE 1: Executive Cover & Telemetry Highlights
    # ══════════════════════════════════════════════════════════════
    s1 = add_base_slide(None, None)
    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(3.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = f"ENTERPRISE PRODUCT SPECIFICATION  ·  {data['category'].upper()}"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = c_cyan
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = data['name']
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = c_white
    p1.space_after = Pt(6)

    p2 = tf1.add_paragraph()
    p2.text = data['tagline'] if data['tagline'] else data['brief']
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = c_gold
    p2.space_after = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = f"Industrial Protocol: {data['protocol']}    |    Domain Architecture: Terrestrial Deterministic Edge Core"
    p3.font.size = Pt(12.5)
    p3.font.color.rgb = c_text
    p3.space_after = Pt(14)

    p4 = tf1.add_paragraph()
    p4.text = data['brief']
    p4.font.size = Pt(13.5)
    p4.font.color.rgb = c_muted

    # 4 Telemetry Metrics on Slide 1
    m_w = Inches(2.76)
    m_h = Inches(1.4)
    m_gap = Inches(0.22)
    m_top = Inches(5.1)
    for i, (lbl, val) in enumerate(data['metrics'][:4]):
        m_left = Inches(0.8) + i * (m_w + m_gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left, m_top, m_w, m_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border_sub
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp1 = ctf.paragraphs[0]
        cp1.text = lbl.upper()
        cp1.alignment = PP_ALIGN.CENTER
        cp1.font.size = Pt(9.5)
        cp1.font.bold = True
        cp1.font.color.rgb = c_muted
        cp1.space_after = Pt(6)

        cp2 = ctf.add_paragraph()
        cp2.text = val
        cp2.alignment = PP_ALIGN.CENTER
        cp2.font.size = Pt(13.5)
        cp2.font.bold = True
        cp2.font.color.rgb = c_cyan

    # ══════════════════════════════════════════════════════════════
    # SLIDE 2: Core Capabilities & Architecture ("What It Delivers")
    # ══════════════════════════════════════════════════════════════
    s2 = add_base_slide(data['name'].upper(), "Core Functional Deliverables & Architectural Pillars")
    
    # 2x2 Grid of Feature Cards
    grid_w = Inches(5.7)
    grid_h = Inches(2.35)
    xs = [Inches(0.8), Inches(6.8)]
    ys = [Inches(1.8), Inches(4.45)]

    for i, (ft_title, ft_desc) in enumerate(data['features'][:4]):
        col = i % 2
        row = i // 2
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xs[col], ys[row], grid_w, grid_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border_sub

        ctf = card.text_frame
        ctf.word_wrap = True
        
        p_num = ctf.paragraphs[0]
        p_num.text = f"PILLAR 0{i+1}"
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = c_cyan
        p_num.space_after = Pt(4)

        p_t = ctf.add_paragraph()
        p_t.text = ft_title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = c_white
        p_t.space_after = Pt(6)

        p_d = ctf.add_paragraph()
        p_d.text = ft_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = c_text

    # ══════════════════════════════════════════════════════════════
    # SLIDE 3: Real-World Operational Use Cases & Field Impact
    # ══════════════════════════════════════════════════════════════
    s3 = add_base_slide(data['name'].upper(), "Field-Validated Operational Use Cases & Quantifiable ROI")
    
    uc_w = Inches(3.72)
    uc_h = Inches(4.8)
    uc_gap = Inches(0.27)
    
    for i, (badge_t, uc_title, uc_desc) in enumerate(data['use_cases'][:3]):
        left_pos = Inches(0.8) + i * (uc_w + uc_gap)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), uc_w, uc_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border_sub

        ctf = card.text_frame
        ctf.word_wrap = True
        
        p_b = ctf.paragraphs[0]
        p_b.text = badge_t.upper()
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = c_gold
        p_b.space_after = Pt(8)

        p_t = ctf.add_paragraph()
        p_t.text = uc_title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = c_white
        p_t.space_after = Pt(12)

        p_d = ctf.add_paragraph()
        p_d.text = uc_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = c_text

    # ══════════════════════════════════════════════════════════════
    # SLIDE 4: Platform vs. Legacy Systems (Competitive Differentiation)
    # ══════════════════════════════════════════════════════════════
    s4 = add_base_slide(data['name'].upper(), "Competitive Advantage: Terrestrial Platform vs. Legacy Systems")

    # Table layout
    col1_w = Inches(5.7)
    col2_w = Inches(5.7)
    table_left = Inches(0.8)
    table_top = Inches(1.8)

    # Column Headers
    h1_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, table_left, table_top, col1_w, Inches(0.55))
    h1_box.fill.solid()
    h1_box.fill.fore_color.rgb = RGBColor(40, 15, 20)
    h1_box.line.color.rgb = c_red
    h1_p = h1_box.text_frame.paragraphs[0]
    h1_p.text = "TRADITIONAL LEGACY SCADA / SILOED SYSTEMS"
    h1_p.font.size = Pt(11)
    h1_p.font.bold = True
    h1_p.font.color.rgb = c_red
    h1_p.alignment = PP_ALIGN.CENTER

    h2_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, table_left + col1_w + Inches(0.3), table_top, col2_w, Inches(0.55))
    h2_box.fill.solid()
    h2_box.fill.fore_color.rgb = RGBColor(10, 45, 60)
    h2_box.line.color.rgb = c_cyan
    h2_p = h2_box.text_frame.paragraphs[0]
    h2_p.text = f"{data['name'].upper()} TERRESTRIAL PLATFORM ADVANTAGE"
    h2_p.font.size = Pt(11)
    h2_p.font.bold = True
    h2_p.font.color.rgb = c_cyan
    h2_p.alignment = PP_ALIGN.CENTER

    row_h = Inches(1.0)
    row_top_start = table_top + Inches(0.65)
    row_gap = Inches(0.12)

    for i, (legacy_pt, terra_pt) in enumerate(data['comparisons'][:4]):
        r_y = row_top_start + i * (row_h + row_gap)
        
        c1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, table_left, r_y, col1_w, row_h)
        c1.fill.solid()
        c1.fill.fore_color.rgb = c_card
        c1.line.color.rgb = c_border_sub
        c1_tf = c1.text_frame
        c1_tf.word_wrap = True
        c1_p = c1_tf.paragraphs[0]
        c1_p.text = legacy_pt
        c1_p.font.size = Pt(10)
        c1_p.font.color.rgb = c_muted

        c2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, table_left + col1_w + Inches(0.3), r_y, col2_w, row_h)
        c2.fill.solid()
        c2.fill.fore_color.rgb = c_card_alt
        c2.line.color.rgb = c_border
        c2_tf = c2.text_frame
        c2_tf.word_wrap = True
        c2_p = c2_tf.paragraphs[0]
        c2_p.text = terra_pt
        c2_p.font.size = Pt(10.5)
        c2_p.font.bold = True
        c2_p.font.color.rgb = c_white

    # ══════════════════════════════════════════════════════════════
    # SLIDE 5: End-to-End System Architecture & Data Pipeline
    # ══════════════════════════════════════════════════════════════
    s5 = add_base_slide(data['name'].upper(), "End-to-End System Architecture & Real-Time Pipeline")

    p_w = Inches(2.76)
    p_h = Inches(4.8)
    p_gap = Inches(0.22)

    for i, (layer_tag, layer_title, layer_desc) in enumerate(data['pipeline'][:4]):
        p_left = Inches(0.8) + i * (p_w + p_gap)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_left, Inches(1.8), p_w, p_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border_sub

        ctf = card.text_frame
        ctf.word_wrap = True
        
        p_tag = ctf.paragraphs[0]
        p_tag.text = layer_tag.upper()
        p_tag.font.size = Pt(9.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = c_cyan
        p_tag.space_after = Pt(8)

        p_t = ctf.add_paragraph()
        p_t.text = layer_title
        p_t.font.size = Pt(13.5)
        p_t.font.bold = True
        p_t.font.color.rgb = c_white
        p_t.space_after = Pt(12)

        p_d = ctf.add_paragraph()
        p_d.text = layer_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = c_text

    # ══════════════════════════════════════════════════════════════
    # SLIDE 6: System Integrator Commercial Playbook & Deployment
    # ══════════════════════════════════════════════════════════════
    s6 = add_base_slide(data['name'].upper(), "System Integrator Deployment Playbook & Commercial Enablement")

    si_w = Inches(3.72)
    si_h = Inches(3.5)
    si_gap = Inches(0.27)

    playbook_steps = [
        ("PHASE 01  ·  CONNECTIVITY & AUDIT", "Zero Rip-and-Replace Onboarding",
         f"Deploy edge container onto existing industrial PC or VM. Connect to plant DCS/PLCs via {data['protocol']}. Ingest and normalize all telemetry in under 48 hours without halting production."),
        ("PHASE 02  ·  TWIN CALIBRATION", "Physics-Model Synchronization",
         f"Feed historical historian data into {data['name']} models. Automatically build directed causal relationship graphs, calibrate thermodynamics/kinematics, and train edge neural operators."),
        ("PHASE 03  ·  AUTONOMOUS GO-LIVE", "Closed-Loop Setpoint Actuation",
         f"Transition from advisory warnings to sub-25ms supervisory actuation. Sync alarm suppressions with plant operators and link work orders into SAP/ERP for verifiable recurring ROI.")
    ]

    for i, (p_tag, p_title, p_desc) in enumerate(playbook_steps):
        s_left = Inches(0.8) + i * (si_w + si_gap)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s_left, Inches(1.8), si_w, si_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = c_border_sub

        ctf = card.text_frame
        ctf.word_wrap = True
        
        pt0 = ctf.paragraphs[0]
        pt0.text = p_tag
        pt0.font.size = Pt(9.5)
        pt0.font.bold = True
        pt0.font.color.rgb = c_gold
        pt0.space_after = Pt(8)

        pt1 = ctf.add_paragraph()
        pt1.text = p_title
        pt1.font.size = Pt(14)
        pt1.font.bold = True
        pt1.font.color.rgb = c_white
        pt1.space_after = Pt(10)

        pt2 = ctf.add_paragraph()
        pt2.text = p_desc
        pt2.font.size = Pt(11)
        pt2.font.color.rgb = c_text

    # Bottom Commercial Callout Box on Slide 6
    callout = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.3))
    callout.fill.solid()
    callout.fill.fore_color.rgb = RGBColor(12, 28, 52)
    callout.line.color.rgb = c_border
    
    ctf_call = callout.text_frame
    ctf_call.word_wrap = True
    ctf_call.vertical_anchor = MSO_ANCHOR.MIDDLE

    cp0 = ctf_call.paragraphs[0]
    cp0.text = f"COMMERCIAL VALUE PROPOSITION FOR SYSTEM INTEGRATORS ({data['name'].upper()})"
    cp0.font.size = Pt(10)
    cp0.font.bold = True
    cp0.font.color.rgb = c_cyan
    cp0.space_after = Pt(4)

    cp1 = ctf_call.add_paragraph()
    cp1.text = "High recurring annual software margins (ARR)  ·  Turnkey deployment in days, not months  ·  Full co-selling & pre-sales engineering support  ·  Certified partner portal at platform.tpcloud.tech"
    cp1.font.size = Pt(11.5)
    cp1.font.bold = True
    cp1.font.color.rgb = c_white

    return prs

# Loop through all 20 applications
generated_files = []

for fname in app_files:
    app_data = parse_app_file(fname)
    clean_name = app_data['name'].replace(' ', '_')
    out_filename = f"{clean_name}_Catalogue.pptx"
    
    prs = generate_pptx_for_app(app_data)
    
    for d in out_dirs:
        out_path = os.path.join(d, out_filename)
        prs.save(out_path)
    
    generated_files.append((app_data['name'], out_filename))
    print(f"Generated {out_filename} successfully for {app_data['name']}!")

print(f"\nAll {len(generated_files)} application catalogues generated successfully in PPTX format!")

