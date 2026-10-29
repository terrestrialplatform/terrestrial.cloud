from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank_layout)

# Background
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(10, 14, 26) # #0a0e1a
bg.line.color.rgb = RGBColor(10, 14, 26)

# Top Cyan Line
top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
top_bar.fill.solid()
top_bar.fill.fore_color.rgb = RGBColor(0, 212, 255)
top_bar.line.fill.background()

# Title box
txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(2.0))
tf = txBox.text_frame
tf.word_wrap = True

p_badge = tf.paragraphs[0]
p_badge.text = "TERRESTRIAL PLATFORM™ | AUTONOMOUS AI & SWARMS"
p_badge.font.size = Pt(11)
p_badge.font.bold = True
p_badge.font.color.rgb = RGBColor(0, 212, 255)
p_badge.space_after = Pt(8)

p_title = tf.add_paragraph()
p_title.text = "TerraAI : Enterprise Product Catalogue"
p_title.font.size = Pt(36)
p_title.font.bold = True
p_title.font.color.rgb = RGBColor(255, 255, 255)
p_title.space_after = Pt(10)

p_sub = tf.add_paragraph()
p_sub.text = "Autonomous Industrial Intelligence & Physics-Informed Predictive Operations Engine"
p_sub.font.size = Pt(16)
p_sub.font.color.rgb = RGBColor(200, 216, 240)

prs.save(r'c:\Users\yashj\Downloads\website\test_deck.pptx')
print("Saved test_deck.pptx successfully!")

